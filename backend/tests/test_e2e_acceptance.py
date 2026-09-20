"""END-TO-END ACCEPTANCE: the complete teacher workflow on real infrastructure.

  register -> course -> resources (PDF, PDF, PPTX) -> exam -> question paper (parsed by AI) -> rubric (AI, edited, approved)
  -> 3 handwritten answer sheets (OCR, identify, map, RAG, evaluate) -> review queue -> OCR correction + re-evaluation
  -> teacher override -> final results -> analytics -> student feedback -> Excel export -> cross-teacher isolation

Uses the real Postgres / Redis / Qdrant / Dramatiq worker / Gemini. Set MARKSWALA_NO_AI_CACHE=1 to bypass the
recorded-response cache and hit the live API for every call.
"""
from __future__ import annotations

import io
import time
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.db.session import SessionLocal
from app.main import create_app
from app.models import (
    Answer,
    AnswerSheet,
    AuditLog,
    Evaluation,
    EvaluationCriterion,
    EvaluationSource,
    StudentResult,
    TeacherOverride,
)
from app.providers.registry import get_vector_store
from tests.conftest import make_client, register
from tests.helpers import FIXTURES, upload, wait_for_job

pytestmark = [pytest.mark.e2e, pytest.mark.live_ai]


def wait_sheets(client, exam_id, status="EVALUATED", timeout=900):
    deadline = time.time() + timeout
    rows = []
    while time.time() < deadline:
        rows = client.get(f"/api/exams/{exam_id}/answer-sheets").json()
        if rows and all(r["status"] in (status, "FAILED") for r in rows):
            return rows
        time.sleep(2)
    raise AssertionError(f"sheets did not reach {status}: {[(r['original_filename'], r['status'], (r['job'] or {}).get('error_message')) for r in rows]}")


def money(x) -> Decimal:
    return Decimal(str(x)).quantize(Decimal("0.01"))


def test_full_workflow(client, worker):
    # ================================================================ 1. authentication
    me = register(client, "prof.rao@university.edu", "Prof. Rao")
    assert client.get("/api/auth/me").json()["email"] == "prof.rao@university.edu"

    # ================================================================ 2. course
    course = client.post("/api/courses", json={"code": "CS201", "name": "Data Structures and Algorithms",
                                               "description": "Second-year core course"}).json()
    cid = course["id"]

    # ================================================================ 3. resources -> chunks -> embeddings -> Qdrant
    resources = []
    for name in ["CS201_Syllabus.pdf", "CS201_Lecture_Notes.pdf", "CS201_Lecture_Slides.pptx"]:
        r = upload(client, f"/api/courses/{cid}/resources", FIXTURES / "course" / name)
        assert r.status_code == 201, r.text
        resources.append(r.json())
    for r in resources:
        assert wait_for_job(client, r["job"]["id"], timeout=300)["status"] == "COMPLETED"
    listed = client.get(f"/api/courses/{cid}/resources").json()
    assert len(listed) == 3 and all(x["status"] == "COMPLETED" and x["chunk_count"] > 0 for x in listed)
    for x in listed:  # the vectors really exist in Qdrant
        assert get_vector_store().count_resource(x["id"]) == x["chunk_count"]

    # ================================================================ 4. exam + question paper (AI parse + arithmetic)
    exam = client.post(f"/api/courses/{cid}/exams", json={"title": "Internal Assessment 1", "pass_percentage": 40}).json()
    eid = exam["id"]
    r = upload(client, f"/api/exams/{eid}/question-paper", FIXTURES / "exam" / "CS201_Internal_Assessment_1.pdf")
    assert r.status_code == 202, r.text
    assert wait_for_job(client, r.json()["paper_job"]["id"], timeout=300)["status"] == "COMPLETED"
    qs = client.get(f"/api/exams/{eid}/questions").json()
    assert len(qs["questions"]) == 5 and qs["computed_total_marks"] == 32
    assert qs["declared_total_marks"] == 32 and qs["total_matches_declared"] is True
    by_no = {q["number"]: q for q in qs["questions"]}
    assert by_no[4]["choice_group"] and by_no[4]["choice_group"] == by_no[5]["choice_group"]
    assert client.post(f"/api/exams/{eid}/questions/confirm").json()["confirmed"] is True

    # ================================================================ 5. rubric: AI draft -> teacher edit -> approve
    gen = client.post(f"/api/exams/{eid}/rubric/generate", json={})
    assert gen.status_code == 202, gen.text
    assert wait_for_job(client, gen.json()["id"], timeout=400)["status"] == "COMPLETED"
    rub = client.get(f"/api/exams/{eid}/rubric").json()
    assert rub["all_balanced"] and rub["version"]["source"] == "AI"
    body = {"units": [{"question_id": u["question_id"], "subquestion_id": u["subquestion_id"],
                       "criteria": [{"id": c["id"], "title": c["title"], "description": c["description"], "expected_points": c["expected_points"], "max_marks": c["max_marks"]}
                                    for c in u["criteria"]]} for u in rub["units"]]}
    body["units"][0]["criteria"][0]["title"] = "Defines a stack (LIFO) correctly"  # the teacher tightens one criterion
    edited = client.put(f"/api/exams/{eid}/rubric/draft", json=body).json()
    assert edited["all_balanced"] and edited["version"]["source"] == "TEACHER"
    approved = client.post(f"/api/exams/{eid}/rubric/draft/approve").json()
    assert approved["version"]["status"] == "APPROVED"
    rubric_version_id = approved["version"]["id"]

    # ================================================================ 6. three handwritten sheets: OCR -> map -> RAG -> evaluate
    files = [("files", (f"student_{k}.pdf", (FIXTURES / "answer_sheets" / f"student_{k}.pdf").read_bytes(), "application/pdf")) for k in "ABC"]
    up = client.post(f"/api/exams/{eid}/answer-sheets", files=files)
    assert up.status_code == 201, up.text
    sheets = wait_sheets(client, eid)
    assert all(s["status"] == "EVALUATED" for s in sheets), [(s["original_filename"], s["status"], (s["job"] or {}).get("error_message")) for s in sheets]
    by_file = {s["original_filename"]: s for s in sheets}
    sheet = {k: by_file[f"student_{k}.pdf"] for k in "ABC"}
    assert {sheet[k]["student"]["roll_number"] for k in "ABC"} == {"CS2024-001", "CS2024-002", "CS2024-003"}

    # ---- results come from persisted evaluations and differ meaningfully between the three students
    res = client.get(f"/api/exams/{eid}/results").json()
    assert res["max_total"] == 32 and len(res["rows"]) == 3
    tot = {r["student"]["roll_number"]: money(r["total"]) for r in res["rows"]}
    A, B, C = tot["CS2024-001"], tot["CS2024-002"], tot["CS2024-003"]
    print(f"\nTOTALS  A={A}  B={B}  C={C}  (max 32)")
    assert A > B > C, "strong > partial > weak must be reflected in the marks"
    assert A >= 22 and C <= 14 and all(0 <= t <= 32 for t in (A, B, C))

    # ---- database-level invariants for every evaluation
    with SessionLocal() as db:
        evs = db.scalars(select(Evaluation).where(Evaluation.is_current)).all()
        assert len(evs) == 3 * 7 - 0  # one per answer: 7 gradable units per sheet
        for ev in evs:
            crits = db.scalars(select(EvaluationCriterion).where(EvaluationCriterion.evaluation_id == ev.id)).all()
            assert ev.ai_total == sum((c.score for c in crits), Decimal(0)), "the total is computed by the backend from criterion scores"
            assert ev.max_total == sum((c.max_score for c in crits), Decimal(0))
            assert all(0 <= c.score <= c.max_score for c in crits) and 0 <= ev.ai_total <= ev.max_total
            assert str(ev.rubric_version_id) == rubric_version_id, "every evaluation is tied to the approved rubric version"
            assert ev.answer_text_snapshot is not None and ev.prompt_version and ev.model
            n_src = db.scalar(select(func.count()).select_from(EvaluationSource).where(EvaluationSource.evaluation_id == ev.id))
            if ev.provider == "system":  # deterministic zero for a blank answer
                assert ev.ai_total == 0 and n_src == 0
            else:
                assert n_src >= 1, "RAG provenance must be stored for AI evaluations"
                assert ev.overall_confidence is not None and ev.evaluation_confidence is not None
        # the internal choice: student A skipped Q5, C skipped Q4 -> each is a zero that does NOT count
        rows = {r["student"]["roll_number"]: {q["label"]: q for q in r["questions"]} for r in res["rows"]}
        assert rows["CS2024-001"]["Q5"]["counted"] is False and rows["CS2024-001"]["Q4"]["counted"] is True
        assert rows["CS2024-003"]["Q4"]["counted"] is False and rows["CS2024-003"]["Q5"]["counted"] is True
        for r in res["rows"]:
            counted = sum(money(q["score"]) for q in r["questions"] if q["counted"])
            assert counted == money(r["total"]), "student total == sum of counted question marks"
        stored = {x.total for x in db.scalars(select(StudentResult)).all()}
        assert stored == {money(r["total"]) for r in res["rows"]}, "persisted results match what the API reports"

    # ================================================================ 7. review queue
    q = client.get(f"/api/exams/{eid}/review?status=ALL").json()
    print(f"REVIEW QUEUE: {len(q['items'])} items, mandatory pending={q['pending_mandatory']}, recommended pending={q['pending_recommended']}")
    assert q["items"], "some answers should be flagged (diagram, unreadable, low confidence, missing answers)"
    for it in q["items"]:
        assert it["reasons"] and it["student"] is None or it["student"]["roll_number"].startswith("CS2024")
    mand_first = [i["mandatory"] for i in q["items"]]
    assert mand_first == sorted(mand_first, reverse=True), "mandatory items are listed first"

    # ================================================================ 8. review one answer: trace, OCR correction, re-evaluation
    detail_B = client.get(f"/api/answer-sheets/{sheet['B']['id']}").json()
    ans_b = {a["label"]: a for a in detail_B["answers"]}["Q2"]
    rv = client.get(f"/api/answers/{ans_b['id']}/review").json()
    ev0 = rv["evaluation"]
    assert rv["question_label"] == "Q2" and rv["student"]["roll_number"] == "CS2024-002" and len(rv["pages"]) >= 1
    assert ev0["attempt"] == 1 and ev0["text_source"] == "ORIGINAL" and len(ev0["sources"]) >= 1 and len(ev0["criteria"]) >= 2
    assert ev0["confidence"]["overall"] is not None and ev0["rubric_version_id"] == rubric_version_id
    assert all(c["evidence"] is not None or c["ai_score"] == 0 for c in ev0["criteria"])
    original_ocr = ans_b["original_text"]

    fix = original_ocr + "\nThe array must be sorted before binary search can be applied, otherwise halving the range is meaningless."
    c1 = client.put(f"/api/answers/{ans_b['id']}/corrected-text", json={"text": fix}).json()
    assert c1["original_text"] == original_ocr, "original OCR is never modified"
    assert c1["corrected_text"] == fix and c1["effective_text"] == fix
    assert client.get(f"/api/answers/{ans_b['id']}/review").json()["evaluation"]["stale"] is True, "the evaluation used the older text"
    rr = client.post(f"/api/answers/{ans_b['id']}/re-evaluate", json={"use_corrected_text": True})
    assert rr.status_code == 202, rr.text
    assert wait_for_job(client, rr.json()["id"], timeout=300)["status"] in ("COMPLETED", "REQUIRES_REVIEW")
    rv2 = client.get(f"/api/answers/{ans_b['id']}/review").json()
    ev1 = rv2["evaluation"]
    assert ev1["attempt"] == 2 and ev1["text_source"] == "CORRECTED" and "must be sorted" in ev1["answer_text_used"] and ev1["stale"] is False
    assert [h["attempt"] for h in rv2["history"]] == [2, 1] and rv2["history"][1]["ai_total"] == ev0["ai_total"], "the first AI result is kept"
    assert rv2["answer"]["original_text"] == original_ocr
    print(f"RE-EVALUATION of B/Q2 with corrected OCR: {ev0['ai_total']} -> {ev1['ai_total']}")

    # ================================================================ 9. teacher override (AI score preserved, audit trail)
    detail_A = client.get(f"/api/answer-sheets/{sheet['A']['id']}").json()
    ans_a = {a["label"]: a for a in detail_A["answers"]}["Q2"]
    rva = client.get(f"/api/answers/{ans_a['id']}/review").json()
    eva = rva["evaluation"]
    crit = next(c for c in eva["criteria"] if float(c["ai_score"]) >= 1)
    before_total = money(next(r for r in client.get(f"/api/exams/{eid}/results").json()["rows"] if r["sheet_id"] == sheet["A"]["id"])["total"])
    ai_score = money(crit["ai_score"])

    assert client.post(f"/api/evaluations/{eva['id']}/override", json={"overrides": [{"evaluation_criterion_id": crit["id"], "teacher_score": float(ai_score) - 1, "reason": " "}]}).status_code == 422
    over = client.post(f"/api/evaluations/{eva['id']}/override", json={"overrides": [{"evaluation_criterion_id": crit["id"], "teacher_score": float(crit["max_score"]) + 1, "reason": "too generous"}]})
    assert over.status_code == 422 and over.json()["error"]["code"] == "score_above_max"
    ok = client.post(f"/api/evaluations/{eva['id']}/override", json={"overrides": [{"evaluation_criterion_id": crit["id"], "teacher_score": float(ai_score) - 1,
                                                                                    "reason": "Derivation skipped a step in the recurrence."}], "notes": "checked page 2"})
    assert ok.status_code == 200, ok.text
    d = ok.json()["evaluation"]
    oc = next(c for c in d["criteria"] if c["id"] == crit["id"])
    assert money(oc["ai_score"]) == ai_score and money(oc["final_score"]) == ai_score - 1, "AI score untouched; final score reflects the teacher"
    assert oc["override"]["reason"] == "Derivation skipped a step in the recurrence." and oc["override"]["teacher_name"] == "Prof. Rao"
    assert money(d["final_total"]) == money(d["ai_total"]) - 1 and d["review"]["status"] == "RESOLVED" and d["review"]["decision"] == "OVERRIDDEN"
    after_total = money(next(r for r in client.get(f"/api/exams/{eid}/results").json()["rows"] if r["sheet_id"] == sheet["A"]["id"])["total"])
    assert after_total == before_total - 1, "final results reflect the override, computed server-side"
    # change again, then revert: history is append-only, AI score still intact
    client.post(f"/api/evaluations/{eva['id']}/override", json={"overrides": [{"evaluation_criterion_id": crit["id"], "teacher_score": float(ai_score) - 0.5, "reason": "Second look: partial credit."}]})
    reverted = client.post(f"/api/evaluations/{eva['id']}/override", json={"revert_criterion_ids": [crit["id"]]}).json()["evaluation"]
    oc = next(c for c in reverted["criteria"] if c["id"] == crit["id"])
    assert oc["override"] is None and money(oc["final_score"]) == ai_score and len(oc["override_history"]) == 2
    # a final override that stays in place for the export
    client.post(f"/api/evaluations/{eva['id']}/override", json={"overrides": [{"evaluation_criterion_id": crit["id"], "teacher_score": float(ai_score) - 1, "reason": "Derivation skipped a step in the recurrence."}]})
    with SessionLocal() as db:
        rows = db.scalars(select(TeacherOverride).where(TeacherOverride.evaluation_id == eva["id"])).all()
        assert len(rows) == 3 and sum(1 for r in rows if r.is_active) == 1
        ec = db.get(EvaluationCriterion, crit["id"])
        assert ec.score == ai_score, "the AI decision is never overwritten"
        actions = [a.action for a in db.scalars(select(AuditLog).order_by(AuditLog.created_at))]
        assert actions.count("evaluation.override") >= 3 and "answer.correct_ocr" in actions and "answer.re_evaluate" in actions

    # a re-evaluation must not silently discard that override
    guard = client.post(f"/api/answers/{ans_a['id']}/re-evaluate", json={"use_corrected_text": False})
    assert guard.status_code == 409 and guard.json()["error"]["code"] == "has_overrides"

    # accept another flagged item outright
    pending = [i for i in client.get(f"/api/exams/{eid}/review?status=PENDING").json()["items"] if i["kind"] == "EVALUATION"]
    if pending:
        acc = client.post(f"/api/evaluations/{pending[0]['evaluation_id']}/accept", json={"notes": "Looks right."})
        assert acc.status_code == 200 and acc.json()["evaluation"]["review"]["decision"] == "ACCEPTED"

    # ================================================================ 10. final results, analytics
    res2 = client.get(f"/api/exams/{eid}/results").json()
    assert {money(r["total"]) for r in res2["rows"]} != {A, B, C}, "results moved after the review actions"
    an = client.get(f"/api/exams/{eid}/analytics").json()
    totals = [r["total"] for r in res2["rows"]]
    assert an["class"]["students"] == 3 and an["class"]["max_total"] == 32
    assert an["class"]["highest"] == max(totals) and an["class"]["lowest"] == min(totals)
    assert abs(an["class"]["average"] - sum(totals) / 3) < 0.01
    assert an["class"]["median"] == sorted(totals)[1]
    assert sum(b["count"] for b in an["class"]["distribution"]) == 3
    assert len(an["questions"]) == 7 and all(q["max_marks"] > 0 for q in an["questions"])
    q2s = next(q for q in an["questions"] if q["label"] == "Q2")
    assert q2s["attempts"] == 3 and q2s["most_missed_criteria"] and 0 <= q2s["percentage"] <= 100
    assert an["topics"] and all(0 <= t["average_percentage"] <= 100 for t in an["topics"])
    assert {s["roll_number"] for s in an["students"]} == {"CS2024-001", "CS2024-002", "CS2024-003"}
    weak_C = next(s for s in an["students"] if s["roll_number"] == "CS2024-003")
    assert weak_C["rubric_failures"], "the weak student's analysis lists the rubric criteria they failed"

    # ================================================================ 11. student feedback from the evaluation record
    fb = client.post(f"/api/answer-sheets/{sheet['C']['id']}/feedback")
    assert fb.status_code == 202, fb.text
    assert wait_for_job(client, fb.json()["id"], timeout=300)["status"] == "COMPLETED"
    result_C = client.get(f"/api/answer-sheets/{sheet['C']['id']}/result").json()
    fbk = result_C["feedback"]
    assert fbk and fbk["narrative"]["improvements"], "a weak student must receive concrete improvements"
    missed = {(m["question"].lower(), m["title"].lower()) for m in fbk["facts"]["missed_criteria"]}
    for imp in fbk["narrative"]["improvements"]:
        assert (imp["question"].lower(), imp["criterion"].lower()) in missed, "feedback cites only criteria the student actually missed"
    assert result_C["feedback_stale"] is False

    # ================================================================ 12. Excel export == database values
    x = client.get(f"/api/exams/{eid}/export.xlsx")
    assert x.status_code == 200 and x.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Final Marks", "Question-wise Marks", "Student Analysis", "Topic Analysis", "Question Analysis", "AI vs Teacher Overrides", "Review Log", "Info"]
    fm = list(wb["Final Marks"].iter_rows(min_row=2, values_only=True))
    api_rows = {r["student"]["roll_number"]: r for r in res2["rows"]}
    assert len(fm) == 3
    for roll, name, total, mx, pct, result, status, pending_reviews in fm:
        r = api_rows[roll]
        assert money(total) == money(r["total"]) and money(mx) == 32 and money(pct) == money(r["percentage"]) and name == r["student"]["full_name"]
        assert result == ("Pass" if r["passed"] else "Fail") and status == r["status"]
    qw = list(wb["Question-wise Marks"].iter_rows(min_row=1, max_row=4, values_only=True))
    assert qw[0][:2] == ("Roll No", "Student") and qw[0][-1] == "Total"
    for row in qw[1:4]:
        r = api_rows[row[0]]
        by_label = {q["label"]: q for q in r["questions"]}
        for head, cell in zip(qw[0][2:-1], row[2:-1]):
            label = head.split(" ")[0]
            expected = money(by_label[label]["score"]) if by_label[label]["counted"] else None
            assert (money(cell) if cell is not None else None) == expected, (row[0], label)
    ov = list(wb["AI vs Teacher Overrides"].iter_rows(min_row=2, values_only=True))
    assert len(ov) == 3 and {o[8] for o in ov} >= {"Derivation skipped a step in the recurrence."} and sum(1 for o in ov if o[10] == "Active") == 1
    assert all(money(o[5]) - money(o[4]) == money(o[7]) for o in ov)
    rl = list(wb["Review Log"].iter_rows(min_row=2, values_only=True))
    assert any(r[7] == "OVERRIDDEN" for r in rl) and any(r[7] == "ACCEPTED" for r in rl or pending == [])
    assert len(list(wb["Question Analysis"].iter_rows(min_row=2, values_only=True))) == 7
    assert len(list(wb["Topic Analysis"].iter_rows(min_row=2, values_only=True))) >= 1

    # ================================================================ 13. isolation: another teacher sees none of it
    intruder = make_client(create_app())
    register(intruder, name="Intruder")
    a_sheet, a_ans, a_ev = sheet["A"]["id"], ans_a["id"], eva["id"]
    for method, url, body in [
        ("get", f"/api/courses/{cid}", None), ("get", f"/api/exams/{eid}", None), ("get", f"/api/exams/{eid}/results", None),
        ("get", f"/api/exams/{eid}/analytics", None), ("get", f"/api/exams/{eid}/export.xlsx", None), ("get", f"/api/exams/{eid}/review", None),
        ("get", f"/api/answer-sheets/{a_sheet}", None), ("get", f"/api/answer-sheets/{a_sheet}/pages/1/image", None), ("get", f"/api/answer-sheets/{a_sheet}/result", None),
        ("get", f"/api/answers/{a_ans}/review", None), ("put", f"/api/answers/{a_ans}/corrected-text", {"text": "x"}),
        ("post", f"/api/evaluations/{a_ev}/override", {"overrides": [{"evaluation_criterion_id": crit["id"], "teacher_score": 0, "reason": "hack"}]}),
        ("post", f"/api/evaluations/{a_ev}/accept", {}), ("post", f"/api/answers/{a_ans}/re-evaluate", {}), ("get", f"/api/resources/{resources[0]['id']}/file", None),
        ("get", f"/api/exams/{eid}/rubric", None), ("get", f"/api/exams/{eid}/question-paper/file", None),
    ]:
        r = getattr(intruder, method)(url, **({"json": body} if body is not None else {}))
        assert r.status_code == 404, f"{method.upper()} {url} leaked: {r.status_code} {r.text[:120]}"
    assert intruder.get("/api/courses").json() == [] and intruder.get("/api/exams").json() == [] and intruder.get("/api/jobs").json() == []

    with SessionLocal() as db:  # nothing was modified by the intruder's attempts
        assert db.scalar(select(func.count()).select_from(TeacherOverride)) == 3
        assert db.scalar(select(func.count()).select_from(AnswerSheet).where(AnswerSheet.deleted_at.is_not(None))) == 0
        assert db.scalar(select(func.count()).select_from(Answer)) == 21
