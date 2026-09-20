"""Text extraction from course materials and chunking with provenance (page / slide / section)."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass

import pymupdf as fitz  # PyMuPDF

from app.providers.base import OCRProvider


@dataclass
class TextUnit:
    text: str
    page: int | None = None
    slide: int | None = None
    section: str | None = None


@dataclass
class Chunk:
    text: str
    page: int | None
    slide: int | None
    section: str | None
    token_count: int


class ExtractionError(Exception):
    """User-actionable extraction failure."""


_HEADING_RE = re.compile(
    r"^(?:(?:module|unit|chapter|section|lecture|topic|part)\s+[\w.\-]+[:.\-\s]|\d+(?:\.\d+)*[.)]?\s+[A-Z])",
    re.IGNORECASE,
)
_BULLET_RE = re.compile(r"^\s*(?:[-*•▪◦]|\d+[.)]|[a-z][.)])\s+")


def is_heading(line: str) -> bool:
    s = line.strip()
    if not s or len(s) > 90 or s.endswith((".", ",", ";")) and not _HEADING_RE.match(s):
        return False
    if _HEADING_RE.match(s):
        return True
    letters = [c for c in s if c.isalpha()]
    return len(letters) >= 4 and s.upper() == s and len(s.split()) <= 10


def _split_by_headings(text: str, current: str | None) -> list[tuple[str | None, str]]:
    """Split a block of text at heading lines -> [(section, body)]."""
    out: list[tuple[str | None, list[str]]] = [(current, [])]
    for line in text.splitlines():
        if is_heading(line):
            out.append((line.strip(), []))
        else:
            out[-1][1].append(line)
    return [(sec, "\n".join(lines).strip()) for sec, lines in out if "\n".join(lines).strip() or sec != current]


# ---------------------------------------------------------------------------------------------------
# extractors
# ---------------------------------------------------------------------------------------------------
def extract_pdf(data: bytes, ocr: OCRProvider | None, on_page=None) -> tuple[list[TextUnit], int]:
    """Text layer first; scanned pages (no text but images) fall back to OCR. Returns (units, ocr_pages)."""
    doc = fitz.open(stream=data, filetype="pdf")
    units: list[TextUnit] = []
    ocr_pages = 0
    section: str | None = None
    try:
        for i, page in enumerate(doc, start=1):
            txt = page.get_text("text").strip()
            if len(txt) < 30 and page.get_images():
                if ocr is None:
                    raise ExtractionError(f"Page {i} is a scan and no OCR provider is configured.")
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
                txt = ocr.transcribe_page(pix.tobytes("png"), page_number=i).text.strip()
                ocr_pages += 1
            for sec, body in _split_by_headings(txt, section):
                section = sec or section
                if body:
                    units.append(TextUnit(body, page=i, section=sec))
            if on_page:
                on_page(i, doc.page_count)
    finally:
        doc.close()
    return units, ocr_pages


def extract_pptx(data: bytes) -> list[TextUnit]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    units: list[TextUnit] = []
    for idx, slide in enumerate(prs.slides, start=1):
        title = None
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip() or None
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                t = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                if t.strip():
                    parts.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            n = slide.notes_slide.notes_text_frame.text.strip()
            if n:
                parts.append("Speaker notes: " + n)
        body = "\n".join(parts).strip()
        if body or title:
            units.append(TextUnit((title + "\n" if title else "") + body, slide=idx, section=title))
    return units


def extract_docx(data: bytes) -> list[TextUnit]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    units: list[TextUnit] = []
    section: str | None = None
    buf: list[str] = []

    def flush():
        if buf:
            units.append(TextUnit("\n".join(buf).strip(), section=section))
            buf.clear()

    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        if p.style is not None and p.style.name.lower().startswith(("heading", "title")):
            flush()
            section = p.text.strip()
        buf.append(p.text.strip())
    for t in doc.tables:
        for row in t.rows:
            buf.append(" | ".join(c.text.strip() for c in row.cells))
    flush()
    return [u for u in units if u.text]


def extract_text(data: bytes) -> list[TextUnit]:
    text = data.decode("utf-8", errors="replace")
    units: list[TextUnit] = []
    for sec, body in _split_by_headings(re.sub(r"^#+\s*", "", text, flags=re.MULTILINE), None):
        if body:
            units.append(TextUnit(body, section=sec))
    return units


def extract_image(data: bytes, ocr: OCRProvider, mime: str) -> list[TextUnit]:
    res = ocr.transcribe_page(data, page_number=1, mime_type=mime)
    return [TextUnit(res.text.strip())] if res.text.strip() else []


# ---------------------------------------------------------------------------------------------------
# chunking
# ---------------------------------------------------------------------------------------------------
def _paragraphs(text: str) -> list[str]:
    """Merge hard-wrapped lines into paragraphs, keeping bullets and blank-line breaks."""
    paras: list[str] = []
    cur: list[str] = []
    for line in text.splitlines():
        s = line.strip()
        if not s:
            if cur:
                paras.append(" ".join(cur))
                cur = []
        elif _BULLET_RE.match(line) or is_heading(line):
            if cur:
                paras.append(" ".join(cur))
            cur = [s]
        else:
            cur.append(s)
    if cur:
        paras.append(" ".join(cur))
    return [re.sub(r"\s+", " ", p).strip() for p in paras if p.strip()]


def _sentences(par: str) -> list[str]:
    return [s for s in re.split(r"(?<=[.!?])\s+(?=[A-Z0-9(\[])", par) if s]


def chunk_units(units: list[TextUnit], *, max_chars: int = 1100, min_chars: int = 200, overlap_chars: int = 120) -> list[Chunk]:
    chunks: list[Chunk] = []
    for u in units:
        pieces: list[str] = []
        for par in _paragraphs(u.text):
            pieces.extend(_sentences(par) if len(par) > max_chars else [par])
        buf = ""
        for piece in pieces:
            if buf and len(buf) + 1 + len(piece) > max_chars:
                chunks.append(Chunk(buf, u.page, u.slide, u.section, len(buf) // 4))
                tail = buf[-overlap_chars:]
                buf = (tail[tail.find(" ") + 1 :] + " " + piece).strip() if len(buf) > overlap_chars else piece
            else:
                buf = f"{buf} {piece}".strip()
        if buf:
            # fold a tiny remainder into the previous chunk of the same unit instead of emitting a stub
            if len(buf) < min_chars and chunks and chunks[-1].page == u.page and chunks[-1].slide == u.slide and chunks[-1].section == u.section:
                prev = chunks[-1]
                chunks[-1] = Chunk(f"{prev.text} {buf}", prev.page, prev.slide, prev.section, (len(prev.text) + len(buf)) // 4)
            else:
                chunks.append(Chunk(buf, u.page, u.slide, u.section, len(buf) // 4))
    return chunks
