import os
import fitz  # PyMuPDF
from pptx import Presentation
import chromadb
from chromadb.config import Settings
from app.core.config import settings
from app.models.resource import Resource
from app.db.session import SessionLocal

# Configure Chroma
chroma_client = chromadb.PersistentClient(path=settings.VECTOR_DB_PATH)

def extract_text_from_pdf(file_path: str) -> str:
    text = ""
    doc = fitz.open(file_path)  # raises if the file is corrupt/not a real PDF
    for page_num in range(len(doc)):
        page = doc.load_page(page_num)
        text += f"\n--- Page {page_num + 1} ---\n"
        text += page.get_text()
    return text

def extract_text_from_pptx(file_path: str) -> str:
    text = ""
    try:
        prs = Presentation(file_path)
        for i, slide in enumerate(prs.slides):
            text += f"\n--- Slide {i + 1} ---\n"
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        run_text = "".join(run.text for run in paragraph.runs)
                        if run_text:
                            text += run_text + "\n"
                if shape.has_table:
                    for row in shape.table.rows:
                        text += " | ".join(cell.text for cell in row.cells) + "\n"
    except Exception as e:
        print(f"Error reading PPTX: {e}")
    return text

def chunk_text(text: str, max_chunk_size: int = 1000) -> list[str]:
    # Simple chunking by paragraphs or max size. MVP implementation.
    chunks = []
    current_chunk = ""
    
    for paragraph in text.split('\n\n'):
        if len(current_chunk) + len(paragraph) > max_chunk_size:
            if current_chunk:
                chunks.append(current_chunk.strip())
            current_chunk = paragraph + "\n\n"
        else:
            current_chunk += paragraph + "\n\n"
            
    if current_chunk:
        chunks.append(current_chunk.strip())
    return chunks

def embed_and_store(resource: Resource, chunks: list[str]):
    collection = chroma_client.get_or_create_collection(name=f"course_{resource.course_id}")
    
    documents = []
    metadatas = []
    ids = []
    
    for i, chunk in enumerate(chunks):
        if not chunk.strip():
            continue
        documents.append(chunk)
        metadatas.append({"resource_id": resource.id, "filename": resource.filename, "course_id": resource.course_id})
        ids.append(f"res_{resource.id}_chunk_{i}")
        
    if documents:
        # Uses Chroma's default local sentence-transformer embedder - fully local, no API calls.
        collection.add(
            documents=documents,
            metadatas=metadatas,
            ids=ids
        )

def process_resource(resource_id: int):
    db = SessionLocal()
    try:
        resource = db.query(Resource).filter(Resource.id == resource_id).first()
        if not resource:
            return

        try:
            text = ""
            if resource.resource_type == "PDF":
                text = extract_text_from_pdf(resource.file_path)
            elif resource.resource_type == "PPTX":
                text = extract_text_from_pptx(resource.file_path)
            elif resource.resource_type == "PPT":
                raise ValueError("Legacy .ppt format is not supported for text extraction, please re-save as .pptx")
            elif resource.resource_type == "TXT":
                with open(resource.file_path, "r", encoding="utf-8", errors="ignore") as f:
                    text = f.read()
            else:
                raise ValueError(f"Unsupported file type for text extraction: {resource.resource_type}")

            chunks = chunk_text(text)
            embed_and_store(resource, chunks)

            resource.status = "COMPLETED"
            resource.chunk_count = len(chunks)
            resource.error_message = None
        except Exception as e:
            print(f"Resource {resource_id} processing failed: {e}")
            resource.status = "FAILED"
            resource.chunk_count = 0
            resource.error_message = str(e)[:500]

        db.commit()
    finally:
        db.close()
