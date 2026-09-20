"""Excel export (openpyxl). Values come from the same functions the API uses, so the workbook matches the database."""
from __future__ import annotations

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    Answer,
    AnswerSheet,
    ConfidenceFlag,
    Course,
    Evaluation,
    EvaluationCriterion,
    Exam,
    Question,
    RubricCriterion,
    Student,
    StudentResult,
    TeacherOverride,
    TeacherReview,
    User,
)
from app.services import analytics_service, results_service
from app.services.marks import units_of

HEAD_FILL = PatternFill("solid", fgColor="C9A7FF")
HEAD_FONT = Font(bold=True, color="172033")
SHEETS = ["Final Marks", "Question-wise Marks", "Student Analysis", "Topic Analysis", "Question Analysis", "AI vs Teacher Overrides", "Review Log", "Info"]


def _table(ws, headers: list[str], rows: list[list], widths: list[int] | None = None) -> None:
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths[i - 1] if widths else max(12, min(48, len(h) + 4)))


def build_workbook(db: Session, exam: Exam) -> bytes:
    course = db.get(Course, exam.course_id)
    analytics = analytics_service.compute(db, exam)
    sheets = list(db.scalars(select(AnswerSheet).where(AnswerSheet.exam_id == exam.id, AnswerSheet.deleted_at.is_(None)).order_by(AnswerSheet.created_at)))
    students = {s.id: s for s in db.scalars(select(Student).where(Student.id.in_([x.student_id for x in sheets if x.student_id])))} if sheets else {}
    results = {s.id: results_service.compute_sheet_result(db, s.id) for s in sheets if s.status.value == "EVALUATED"}
    questions = list(db.scalars(select(Question).where(Question.exam_id == exam.id).order_by(Question.position)))
    units = units_of(questions)

    def ident(sh: AnswerSheet) -> tuple[str, str]:
        st = students.get(sh.student_id) if sh.student_id else None
        return (st.roll_number if st else "(unidentified)", st.full_name if st else sh.original_filename)

    wb = Workbook()
    wb.remove(wb.active)
    ordered = sorted(results.values(), key=lambda r: (ident(next(s for s in sheets if s.id == r.sheet_id))[0]))
    by_sheet = {s.id: s for s in sheets}

    # 1. Final Marks
    ws = wb.create_sheet("Final Marks")
    rows = []
    for r in ordered:
        roll, name = ident(by_sheet[r.sheet_id])
        rows.append([roll, name, float(r.total), float(r.max_total), float(r.percentage), "Pass" if r.passed else "Fail", r.status.value, r.pending_reviews + r.sheet_flags_open])
    _table(ws, ["Roll No", "Student", "Total", "Maximum", "Percentage", "Result", "Status", "Reviews pending"], rows, [16, 28, 10, 10, 12, 10, 14, 16])
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        row[0].number_format = "0.00"

    # 2. Question-wise Marks
    ws = wb.create_sheet("Question-wise Marks")
    heads = ["Roll No", "Student"] + [f"{u.label} (/{float(u.max_marks):g})" for u in units] + ["Total"]
    rows = []
    for r in ordered:
        roll, name = ident(by_sheet[r.sheet_id])
        cells = []
        for us in r.units:
            cells.append(float(us.final_score) if (us.counted and us.evaluated) else None)  # None = alternative not counted / not evaluated
        rows.append([roll, name, *cells, float(r.total)])
    _table(ws, heads, rows)
    ws.cell(row=len(rows) + 3, column=1, value="Blank = an internal-choice alternative that does not count towards the total, or not evaluated.")

    # 3. Student Analysis
    ws = wb.create_sheet("Student Analysis")
    rows = []
    fb = {x.answer_sheet_id: x.feedback for x in db.scalars(select(StudentResult).where(StudentResult.exam_id == exam.id))}
    for s in analytics["students"]:
        f = fb.get(uuid_or_none(s["sheet_id"])) or {}
        rows.append([s["roll_number"], s["name"], s["total"], s["max_total"], s["percentage"], "; ".join(s["weak_topics"]), "; ".join(s["strong_topics"]),
                     len(s["rubric_failures"]), "; ".join(f"{x['question']}: {x['criterion']}" for x in s["rubric_failures"]),
                     (f.get("narrative") or {}).get("summary", "")])
    _table(ws, ["Roll No", "Student", "Total", "Maximum", "Percentage", "Weak topics", "Strong topics", "Criteria missed", "Missed criteria (question: criterion)", "Feedback summary"], rows,
           [16, 26, 9, 9, 11, 34, 34, 12, 60, 70])

    # 4. Topic Analysis
    ws = wb.create_sheet("Topic Analysis")
    _table(ws, ["Topic", "Average %", "Students", "Class level"], [[t["topic"], t["average_percentage"], t["students"], t["class"]] for t in analytics["topics"]], [42, 12, 10, 12])

    # 5. Question Analysis
    ws = wb.create_sheet("Question Analysis")
    rows = []
    for q in analytics["questions"]:
        mm = q["most_missed_criteria"][0] if q["most_missed_criteria"] else None
        rows.append([q["label"], q["topic"] or "", q["max_marks"], q["attempts"], q["students"], q["average"], q["percentage"], q["highest"], q["lowest"],
                     f"{mm['criterion']} ({mm['percentage']}%)" if mm else ""])
    _table(ws, ["Question", "Topic", "Max marks", "Attempted by", "Students", "Average", "% achieved", "Highest", "Lowest", "Most missed criterion"], rows, [12, 34, 10, 13, 10, 10, 11, 9, 9, 60])

    # 6. AI vs Teacher Overrides
    ws = wb.create_sheet("AI vs Teacher Overrides")
    ov_rows = db.execute(
        select(TeacherOverride, EvaluationCriterion, Evaluation, Answer, User)
        .join(EvaluationCriterion, EvaluationCriterion.id == TeacherOverride.evaluation_criterion_id)
        .join(Evaluation, Evaluation.id == TeacherOverride.evaluation_id).join(Answer, Answer.id == Evaluation.answer_id)
        .join(User, User.id == TeacherOverride.teacher_id)
        .where(Answer.answer_sheet_id.in_([s.id for s in sheets]) if sheets else False)
        .order_by(TeacherOverride.created_at)
    ).all() if sheets else []
    unit_label = {(u.question_id, u.subquestion_id): u.label for u in units}
    rc = {c.id: c.title for c in db.scalars(select(RubricCriterion).where(RubricCriterion.id.in_([e[1].rubric_criterion_id for e in ov_rows])))} if ov_rows else {}
    rows = []
    for o, ec, ev, ans, teacher in ov_rows:
        roll, name = ident(by_sheet[ans.answer_sheet_id])
        rows.append([roll, name, unit_label.get((ans.question_id, ans.subquestion_id), ""), rc.get(ec.rubric_criterion_id, ""), float(o.ai_score), float(o.teacher_score),
                     float(o.final_score), float(o.teacher_score - o.ai_score), o.reason, teacher.full_name, "Active" if o.is_active else "Superseded",
                     o.created_at.replace(tzinfo=None) if o.created_at else None])
    _table(ws, ["Roll No", "Student", "Question", "Criterion", "AI score", "Teacher score", "Final score", "Difference", "Reason", "Teacher", "State", "When (UTC)"], rows,
           [16, 24, 11, 40, 10, 13, 11, 11, 50, 22, 12, 20])

    # 7. Review Log
    ws = wb.create_sheet("Review Log")
    rv_rows = db.execute(
        select(TeacherReview, Evaluation, Answer).join(Evaluation, Evaluation.id == TeacherReview.evaluation_id).join(Answer, Answer.id == Evaluation.answer_id)
        .where(TeacherReview.answer_sheet_id.in_([s.id for s in sheets])).order_by(TeacherReview.created_at)
    ).all() if sheets else []
    names = {u.id: u.full_name for u in db.scalars(select(User).where(User.id.in_([r[0].reviewer_id for r in rv_rows if r[0].reviewer_id])))} if rv_rows else {}
    flags: dict = {}
    for f in db.scalars(select(ConfidenceFlag).where(ConfidenceFlag.evaluation_id.in_([r[1].id for r in rv_rows]))) if rv_rows else []:
        flags.setdefault(f.evaluation_id, []).append(f.detail)
    rows = []
    for rv, ev, ans in rv_rows:
        roll, name = ident(by_sheet[ans.answer_sheet_id])
        rows.append([roll, name, unit_label.get((ans.question_id, ans.subquestion_id), ""), "Mandatory" if rv.mandatory else "Recommended", " | ".join(flags.get(ev.id, [])),
                     round(ev.overall_confidence or 0, 3), rv.status.value, rv.decision.value if rv.decision else "", names.get(rv.reviewer_id, ""), rv.notes or "",
                     rv.resolved_at.replace(tzinfo=None) if rv.resolved_at else None])
    _table(ws, ["Roll No", "Student", "Question", "Review type", "Why flagged", "Overall confidence", "Status", "Decision", "Reviewer", "Notes", "Resolved (UTC)"], rows,
           [16, 24, 11, 13, 70, 12, 11, 12, 22, 40, 20])

    # 8. Info
    ws = wb.create_sheet("Info")
    for row in [["Exam", exam.title], ["Course", f"{course.code} — {course.name}"], ["Students evaluated", len(results)],
                ["Exported (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")], ["Generated by", "MarksWala"],
                ["Note", "Marks are computed from stored evaluations and teacher overrides; AI scores are never overwritten."]]:
        ws.append(row)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 90
    for c in ws["A"]:
        c.font = Font(bold=True)

    assert wb.sheetnames == SHEETS
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def uuid_or_none(v: str):
    import uuid

    try:
        return uuid.UUID(v)
    except (ValueError, TypeError):
        return None
