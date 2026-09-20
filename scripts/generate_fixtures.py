"""Generate the realistic DSA test fixtures under ./fixtures.

    backend/.venv/Scripts/python scripts/generate_fixtures.py

Outputs
  fixtures/course/CS201_Syllabus.pdf            printed syllabus
  fixtures/course/CS201_Lecture_Notes.pdf       multi-page lecture notes (module headings)
  fixtures/course/CS201_Lecture_Slides.pptx     10-slide deck
  fixtures/exam/CS201_Internal_Assessment_1.pdf question paper (32 marks, internal choice Q4/Q5)
  fixtures/answer_sheets/student_{A,B,C}.pdf    handwriting-style scans (Strong / Partial / Weak)

Handwriting is synthesised from open-licence handwriting fonts (Google Fonts, downloaded on first run)
with ruled paper, ink jitter, skew and scan noise. It is a realistic *test stimulus*, not real student handwriting.
"""
from __future__ import annotations

import random
import re
import sys
import urllib.request
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

sys.path.insert(0, str(Path(__file__).parent))
import fixture_content as C  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "fixtures"
FONT_DIR = ROOT / "scripts" / ".fonts"
FONT_URLS = {
    "IndieFlower-Regular.ttf": "ofl/indieflower/IndieFlower-Regular.ttf",
    "ShadowsIntoLight.ttf": "ofl/shadowsintolight/ShadowsIntoLight.ttf",
    "GochiHand-Regular.ttf": "ofl/gochihand/GochiHand-Regular.ttf",
}


# --------------------------------------------------------------------------------------------- printed PDFs
def _styles():
    ss = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("t", parent=ss["Title"], fontSize=20, spaceAfter=12),
        "h": ParagraphStyle("h", parent=ss["Heading2"], fontSize=14, spaceBefore=14, spaceAfter=6),
        "p": ParagraphStyle("p", parent=ss["BodyText"], fontSize=11, leading=15, spaceAfter=8),
        "q": ParagraphStyle("q", parent=ss["BodyText"], fontSize=12, leading=17, spaceAfter=10),
        "c": ParagraphStyle("c", parent=ss["BodyText"], fontSize=11, alignment=1, spaceAfter=8),
    }


def _pdf(path: Path, story) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm, topMargin=20 * mm, bottomMargin=20 * mm,
                      title=path.stem, author="MarksWala fixtures").build(story)


def make_syllabus() -> None:
    s = _styles()
    story = [Paragraph(f"{C.COURSE_CODE} — {C.COURSE_NAME}: Syllabus", s["title"])]
    for h, paras in C.SYLLABUS:
        story.append(Paragraph(h, s["h"]))
        story += [Paragraph(p, s["p"]) for p in paras]
    _pdf(OUT / "course" / "CS201_Syllabus.pdf", story)


def make_notes() -> None:
    s = _styles()
    story = [Paragraph(f"{C.COURSE_NAME} — Lecture Notes", s["title"])]
    for i, (h, paras) in enumerate(C.NOTES):
        if i and i % 2 == 0:
            story.append(PageBreak())
        story.append(Paragraph(h, s["h"]))
        story += [Paragraph(p, s["p"]) for p in paras]
    _pdf(OUT / "course" / "CS201_Lecture_Notes.pdf", story)


def make_slides() -> None:
    prs = Presentation()
    for i, (title, bullets) in enumerate(C.SLIDES):
        layout = prs.slide_layouts[0 if i == 0 else 1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        body = slide.placeholders[1]
        tf = body.text_frame
        tf.text = bullets[0]
        for b in bullets[1:]:
            tf.add_paragraph().text = b
        if i:
            for p in tf.paragraphs:
                p.font.size = Pt(24)
    path = OUT / "course" / "CS201_Lecture_Slides.pptx"
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


def make_question_paper() -> None:
    s = _styles()
    story = [Paragraph(C.PAPER_TITLE, s["title"])]
    story += [Paragraph(l, s["c"]) for l in C.PAPER_HEADER]
    story.append(Spacer(1, 8))
    for section, items in C.PAPER:
        story.append(Paragraph(f"<b>{section}</b>", s["h"]))
        for label, text, marks in items:
            if label == "OR":
                story.append(Paragraph("<b>— OR —</b>", s["c"]))
            elif text is None:
                story.append(Paragraph(f"<b>{label}.</b>", s["q"]))
            else:
                shown = label.replace("(", " (", 1) if "(" in label else label
                story.append(Paragraph(f"<b>{shown}</b>&nbsp;&nbsp;{text}&nbsp;&nbsp;<b>[{marks} marks]</b>", s["q"]))
    _pdf(OUT / "exam" / "CS201_Internal_Assessment_1.pdf", story)


# ------------------------------------------------------------------------------------------ handwriting scans
def _fonts() -> None:
    FONT_DIR.mkdir(parents=True, exist_ok=True)
    for name, rel in FONT_URLS.items():
        dest = FONT_DIR / name
        if not dest.exists():
            print("downloading font", name)
            urllib.request.urlretrieve(f"https://github.com/google/fonts/raw/main/{rel}", dest)


W, H = 1240, 1754  # A4 @150 dpi
LINE_H = 50
MARGIN_L, MARGIN_R, TOP = 150, 90, 250


def _label_for(style: str, label: str) -> str:
    m = re.match(r"Q(\d+)(?:\(([a-z])\))?$", label)
    n, sub = m.group(1), m.group(2)
    if style == "A":
        return label
    if style == "B":
        return f"Ans {n}({sub})." if sub else f"Ans {n}."
    return f"Q.{n} {sub})" if sub else f"Q.{n}"


def _wrap(draw, text, font, max_w):
    lines = []
    for raw in text.split("\n"):
        if raw.startswith("[[DIAGRAM"):
            lines.append(raw)
            continue
        words, cur = raw.split(), ""
        for w in words:
            t = (cur + " " + w).strip()
            if draw.textlength(t, font=font) <= max_w:
                cur = t
            else:
                lines.append(cur)
                cur = w
        lines.append(cur)
    return lines


def _paper(rng) -> Image.Image:
    img = Image.new("RGB", (W, H), (250, 248, 240))
    d = ImageDraw.Draw(img)
    for y in range(TOP - 8, H - 100, LINE_H):
        d.line([(60, y), (W - 60, y)], fill=(178, 198, 226), width=2)
    d.line([(MARGIN_L - 25, 0), (MARGIN_L - 25, H)], fill=(226, 130, 130), width=3)
    return img


def _ink(rng):
    return (rng.randint(20, 45), rng.randint(30, 55), rng.randint(95, 140))


def _draw_line(page, text, x, y, font, rng, size_jitter=True):
    layer = Image.new("RGBA", (W, LINE_H + 40), (0, 0, 0, 0))
    ld = ImageDraw.Draw(layer)
    cx = 0
    for word in text.split(" "):
        ld.text((cx, 6 + rng.randint(-2, 2)), word, font=font, fill=_ink(rng) + (235,))
        cx += ld.textlength(word + " ", font=font) + rng.randint(0, 4)
    layer = layer.rotate(rng.uniform(-0.5, 0.5), resample=Image.BICUBIC, expand=False)
    page.paste(layer, (int(x), int(y - 14)), layer)


def _stack_diagram(page, x, y, rng):
    d = ImageDraw.Draw(page)
    ink = _ink(rng)
    bx = x + 40
    for i in range(3):  # three boxes
        top = y + 15 + i * 42
        d.rectangle([bx, top, bx + 180, top + 40], outline=ink, width=3)
    d.line([bx, y + 10, bx, y + 145], fill=ink, width=4)
    d.line([bx + 180, y + 10, bx + 180, y + 145], fill=ink, width=4)
    d.line([bx + 260, y + 80, bx + 200, y + 30], fill=ink, width=3)
    d.line([bx + 200, y + 30, bx + 214, y + 46], fill=ink, width=3)
    d.line([bx + 200, y + 30, bx + 224, y + 30], fill=ink, width=3)


def make_sheet(key: str) -> None:
    st = C.STUDENTS[key]
    rng = random.Random(f"markswala-{key}")
    font = ImageFont.truetype(str(FONT_DIR / st["font"]), 36)
    small = ImageFont.truetype(str(FONT_DIR / st["font"]), 30)
    label_font = ImageFont.truetype(str(FONT_DIR / st["font"]), 40)
    probe = ImageDraw.Draw(Image.new("RGB", (10, 10)))
    max_w = W - MARGIN_L - MARGIN_R

    # 1) lay out all lines: ("label"|"text"|"diagram", str)
    rows: list[tuple[str, str]] = []
    for label, text in st["answers"]:
        rows.append(("label", _label_for(key, label)))
        for ln in _wrap(probe, text, font, max_w):
            rows.append(("diagram", "") if ln.startswith("[[DIAGRAM") else ("text", ln))
        rows.append(("gap", ""))

    # 2) paginate
    pages: list[Image.Image] = []
    page, y = _paper(rng), TOP
    first = True

    def header(pg):
        _draw_line(pg, f"Name: {st['name']}", 150, 90, small, rng)
        _draw_line(pg, f"Roll No: {st['roll']}", 700, 90, small, rng)
        _draw_line(pg, "Subject: Data Structures and Algorithms (CS201)   Internal Assessment 1", 150, 145, small, rng)

    header(page)
    for kind, val in rows:
        need = 4 * LINE_H if kind == "diagram" else LINE_H
        if y + need > H - 110:
            _draw_line(page, f"- {len(pages) + 1} -", W // 2 - 20, H - 70, small, rng)
            pages.append(page)
            page, y = _paper(rng), TOP - 60
        if kind == "label":
            _draw_line(page, val, MARGIN_L, y, label_font, rng)
            ImageDraw.Draw(page).line([(MARGIN_L, y + 38), (MARGIN_L + probe.textlength(val, font=label_font), y + 38)], fill=_ink(rng), width=2)
            y += LINE_H
        elif kind == "text":
            _draw_line(page, val, MARGIN_L + rng.randint(0, 6), y, font, rng)
            y += LINE_H
        elif kind == "diagram":
            _stack_diagram(page, MARGIN_L, y, rng)
            y += 4 * LINE_H
        else:
            y += LINE_H // 2
    _draw_line(page, f"- {len(pages) + 1} -", W // 2 - 20, H - 70, small, rng)
    pages.append(page)

    # 3) scan effects
    final = []
    for pg in pages:
        pg = pg.rotate(rng.uniform(-0.6, 0.6), resample=Image.BICUBIC, fillcolor=(250, 248, 240)).filter(ImageFilter.GaussianBlur(0.7))
        px = pg.load()
        for _ in range(9000):  # sensor noise
            x, yy = rng.randrange(W), rng.randrange(H)
            v = rng.randint(-14, 14)
            r, g, b = px[x, yy]
            px[x, yy] = (max(0, min(255, r + v)), max(0, min(255, g + v)), max(0, min(255, b + v)))
        final.append(pg)
    out = OUT / "answer_sheets" / f"student_{key}.pdf"
    out.parent.mkdir(parents=True, exist_ok=True)
    final[0].save(out, "PDF", resolution=150, save_all=True, append_images=final[1:], quality=80)
    print(f"student_{key}.pdf: {len(final)} pages, {out.stat().st_size // 1024} KB")


def main() -> None:
    _fonts()
    make_syllabus()
    make_notes()
    make_slides()
    make_question_paper()
    for k in C.STUDENTS:
        make_sheet(k)
    print("fixtures written to", OUT)


if __name__ == "__main__":
    main()
